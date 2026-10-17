import copy, re
import docx
from docx.shared import Pt, RGBColor, Emu
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = "/Users/jose/Documents/SEM2/Ciberseguridad/"
TEMPLATE = BASE + "Guia_Docente_Clase_5 - 2026-2.docx"
USABLE_CM = 16.59


def new_doc():
    d = docx.Document(TEMPLATE)
    body = d.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    return d


def _run(par, text, size=10.5, bold=False, color=None, italic=False, font="Calibri"):
    r = par.add_run(text)
    r.font.name = font
    r._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def rich(par, text, size=10.5, color=None):
    """**negrita** y `codigo` dentro del texto."""
    for tok in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            _run(par, tok[2:-2], size, bold=True, color=color)
        elif tok.startswith("`"):
            _run(par, tok[1:-1], size - 0.5, font="Consolas", color="7F1010")
        else:
            _run(par, tok, size, color=color)


def banner(d, titulo, subtitulo):
    p = d.add_paragraph()
    _run(p, "UNIVERSIDAD DE MEDELLÍN · FACULTAD DE INGENIERÍA · INGENIERÍA DE SISTEMAS", 8.5, True, "595959")
    p.paragraph_format.space_after = Pt(2)
    p = d.add_paragraph()
    _run(p, titulo, 20, True, "C00000")
    p.paragraph_format.space_after = Pt(2)
    p = d.add_paragraph()
    _run(p, subtitulo, 13, True, "333333")
    p.paragraph_format.space_after = Pt(12)


def h1(d, t):
    d.add_paragraph(t, style="Heading 1")


def h2(d, t):
    d.add_paragraph(t, style="Heading 2")


def para(d, t, size=10.5):
    p = d.add_paragraph()
    rich(p, t, size)
    p.paragraph_format.space_after = Pt(6)
    return p


def bullets(d, items):
    for t in items:
        p = d.add_paragraph(style="List Bullet")
        rich(p, t)


def code(d, text):
    for line in text.strip("\n").split("\n"):
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Emu(180000)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F2F2F2")
        pPr.append(shd)
        _run(p, line if line else " ", 8.5, font="Consolas")
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def table(d, widths_cm, rows, header=True, size=9.5):
    t = d.add_table(rows=len(rows), cols=len(widths_cm))
    t.style = "Table Grid"
    t.autofit = False
    tblPr = t._tbl.tblPr
    for g in t._tbl.findall(qn("w:tblGrid")):
        t._tbl.remove(g)
    for tag in ("w:tblW", "w:tblLayout"):
        for e in tblPr.findall(qn(tag)):
            tblPr.remove(e)
    tot = sum(int(w * 567) for w in widths_cm)
    tw = OxmlElement("w:tblW"); tw.set(qn("w:type"), "dxa"); tw.set(qn("w:w"), str(tot)); tblPr.append(tw)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    grid = OxmlElement("w:tblGrid")
    for w in widths_cm:
        gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), str(int(w * 567))); grid.append(gc)
    t._tbl.insert(list(t._tbl).index(tblPr) + 1, grid)
    for i, row in enumerate(rows):
        for j, txt in enumerate(row):
            c = t.cell(i, j)
            c.width = Emu(int(widths_cm[j] * 360000))
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            if header and i == 0:
                _run(p, txt, size, True, "C00000")
                tcPr = c._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F8E8E8")
                tcPr.append(shd)
            else:
                rich(p, txt, size)
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


REFS = [
    "Brezinski, D. y Killalea, T. (2002). *Guidelines for Evidence Collection and Archiving* (RFC 3227, BCP 55). IETF. https://www.rfc-editor.org/rfc/rfc3227",
    "ISO/IEC 27037:2012. *Guidelines for identification, collection, acquisition and preservation of digital evidence*. ISO. https://www.iso.org/standard/44381.html (edición vigente a oct-2026).",
    "Kent, K., Chevalier, S., Grance, T. y Dang, H. (2006). *Guide to Integrating Forensic Techniques into Incident Response* (NIST SP 800-86). https://csrc.nist.gov/pubs/sp/800/86/final",
    "Nelson, A., Rekhi, S., Souppaya, M. y Scarfone, K. (2025). *Incident Response Recommendations and Considerations for Cybersecurity Risk Management: A CSF 2.0 Community Profile* (NIST SP 800-61 Rev. 3). https://csrc.nist.gov/pubs/sp/800/61/r3/final",
    "SWGDE (2026). *Best Practices for Computer Forensic Acquisition* (17-F-002-2.1). https://www.swgde.org/documents/published-complete-listing/17-f-002-2-1/",
    "Volatility Foundation. *Volatility 3 documentation* y *Linux Tutorial* (plugins `banners.Banners`, `linux.pslist`, `linux.psaux`, `linux.lsof`, `linux.sockstat`, `linux.bash`). https://volatility3.readthedocs.io/en/latest/getting-started-linux-tutorial.html",
    "Microsoft. *AVML: Acquire Volatile Memory for Linux* (herramienta de adquisición en espacio de usuario). https://github.com/microsoft/avml; y Volatility Foundation, *dwarf2json* y repositorio comunitario de símbolos ISF https://github.com/Abyss-W4tcher/volatility3-symbols",
    "Debian Project. *Debian 13 «trixie» Release Information* (versión estable vigente a octubre de 2026: 13.7). https://www.debian.org/releases/trixie/",
    "Congreso de la República de Colombia (2004). *Ley 906 de 2004, Código de Procedimiento Penal*, arts. 254 y ss. (cadena de custodia); y Fiscalía General de la Nación, *Manual del Sistema de Cadena de Custodia*. https://www.fiscalia.gov.co/",
    "SANS Institute. *FOR508: Advanced Incident Response, Threat Hunting, and Digital Forensics* (referencia de práctica de la industria). https://www.sans.org/cyber-security-courses/advanced-incident-response-threat-hunting-training/",
]


def refs(d):
    for r in REFS:
        p = d.add_paragraph(style="List Number")
        rich(p, r.replace("*", ""), 9.5)
